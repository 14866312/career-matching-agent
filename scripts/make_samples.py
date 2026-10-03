'''Generate the fictional resumes and student fixtures used by tests and demos.'''

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = ROOT / 'samples'
RESUME_LINES = (
    '虚构学生简历 · 仅用于功能演示',
    '专业：软件工程',
    '职业意向：Java开发工程师；城市：北京',
    '技能：了解Java，在课程项目中编写了图书借阅类和单元测试。',
    '技能：了解SQL，在课程项目中编写了查询与联表语句。',
    '技能：使用Git提交课程代码并处理过一次合并冲突。',
    '项目经历：与三位同学协作开发图书管理课程项目，负责Java借阅逻辑。',
    '通用素质：在团队协作中记录分工，每周向同学沟通问题与进度。',
    '证书：暂无个人证书。',
    '声明：以上经历完全虚构，不包含真实姓名、联系方式或身份信息。',
)
STUDENT_FILENAMES = {
    'zero': '学生-零技能.json',
    'partial': '学生-部分匹配.json',
    'pending': '学生-待确认.json',
    'related': '学生-关联技能.json',
    'levels': '学生-等级案例.json',
}


def ability(tag, label, level, evidence, confirmed=True):
    return {
        'tag_id': tag,
        'label': label,
        'level': level,
        'evidence': evidence,
        'confirmed': confirmed,
    }


def student_profiles():
    base = {
        'major': '软件工程',
        'skills': [],
        'certificates': [],
        'qualities': [],
        'experiences': '虚构课程项目，仅用于演示',
        'intention': {'target_job_id': 'java', 'city': ''},
        'confirmed': True,
        'advantages': [],
        'improvements': [],
    }
    partial = {
        **base,
        'skills': [
            ability('java', 'Java', 1, 'Java课程项目：编写图书借阅类'),
            ability('sql', 'SQL', 1, 'SQL课程项目：查询与联表练习'),
        ],
    }
    return {
        'zero': base,
        'partial': partial,
        'pending': {
            **base,
            'skills': [ability('java', 'Java', 2, '', False)],
        },
        'related': {
            **base,
            'skills': [ability('javascript', 'JavaScript', 2, 'JavaScript课程练习')],
        },
        'levels': {
            **base,
            'skills': [
                ability('java', 'Java', 3, '独立实现Java课程服务'),
                ability('sql', 'SQL', 1, 'SQL查询课程作业'),
            ],
        },
    }


def write_docx(path: Path) -> None:
    document = Document()
    document.core_properties.author = 'Student Career Planning'
    document.core_properties.title = '虚构学生演示简历'
    normal = document.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(11)

    title = document.add_heading(RESUME_LINES[0], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in RESUME_LINES[1:]:
        document.add_paragraph(line)
    document.save(path)


def build_samples(output_dir: Path | str = DEFAULT_OUTPUT_DIR, *, include_pdf: bool = False):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    text_path = output_dir / '虚构简历.txt'
    docx_path = output_dir / '虚构简历.docx'
    text_path.write_text('\n'.join(RESUME_LINES) + '\n', encoding='utf-8')
    write_docx(docx_path)

    output = {'txt': text_path, 'docx': docx_path}
    for key, profile in student_profiles().items():
        path = output_dir / STUDENT_FILENAMES[key]
        path.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        output[key] = path

    if include_pdf:
        if __package__:  # Imported as scripts.make_samples.
            from .make_sample_pdf import build_pdf
        else:  # Executed directly as python scripts/make_samples.py.
            from make_sample_pdf import build_pdf

        pdf_path = output_dir / '虚构简历.pdf'
        build_pdf(text_path, pdf_path)
        output['pdf'] = pdf_path

    return output


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument('--include-pdf', action='store_true')
    return parser.parse_args(argv)


def main(argv=None) -> None:
    args = parse_args(argv)
    build_samples(args.output_dir, include_pdf=args.include_pdf)
    print(f'Fictional samples generated in {args.output_dir}.')


if __name__ == '__main__':
    main()
