"""HTTP contract and privacy regression tests. AI functions are explicitly stubbed."""
import pytest
from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from backend.app import main
from backend.app.models import StudentProfile
from backend.app.resume import MAX_BYTES


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(main, 'ROOT', tmp_path)
    return TestClient(main.app, base_url='http://127.0.0.1')


def student():
    return {'major': '软件工程', 'confirmed': True, 'skills': [{'tag_id': 'java', 'label': 'Java', 'confirmed': True, 'level': 2, 'evidence': 'Java课程项目'}], 'intention': {'target_job_id': 'java', 'city': '北京'}}


def test_docx_upload_passes_header_text_to_resume_model(client, monkeypatch):
    from backend.app import llm

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = '姓名：虚构同学。专业：软件工程。'
    doc.add_paragraph('使用Java完成课程项目。')
    output = BytesIO()
    doc.save(output)

    async def model(_instruction, payload, **_options):
        assert '姓名：虚构同学。专业：软件工程。' in payload['resume_text']
        return {'name': '虚构同学', 'major': '软件工程',
                'experiences': '使用Java完成课程项目。',
                'skills': [{'tag_id': 'java', 'evidence': '使用Java完成课程项目'}]}

    monkeypatch.setattr(llm, 'call_json', model)
    response = client.post('/api/resume/parse', files={
        'file': ('resume.docx', output.getvalue(),
                 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    })
    assert response.status_code == 200
    result = response.json()
    assert result['name'] == '虚构同学'
    assert result['profile']['major'] == '软件工程'
    assert result['profile']['skills'][0]['tag_id'] == 'java'
    assert '虚构同学' not in str(result['profile'])


@pytest.mark.parametrize('route', ['/api/matches', '/api/recommendations', '/api/reports'])
def test_confirmation_is_not_a_gate(client, route, monkeypatch):
    candidate = student()
    candidate['confirmed'] = False
    payload = {'student': candidate}
    if route != '/api/recommendations':
        payload['job_id'] = 'java'
    if route == '/api/reports':
        async def advice(*_args):
            return {'fit_evaluation': '按当前资料核对', 'learning_directions': [], 'learning_steps': []}
        monkeypatch.setattr(main, 'generate_advice', advice)

    response = client.post(route, json=payload)
    assert response.status_code == 200
    assert 'error' not in response.json()


def test_positive_skill_counts_without_confirmation_or_evidence(client):
    candidate = student()
    candidate['skills'][0]['confirmed'] = False
    unadopted = client.post('/api/matches', json={'job_id': 'java', 'student': candidate})
    assert unadopted.status_code == 200
    java = next(item for item in unadopted.json()['items'] if item['tag_id'] == 'java')
    assert java['status'] == 'satisfied'
    assert java in unadopted.json()['satisfied_items']

    candidate['skills'][0]['confirmed'] = True
    candidate['skills'][0]['evidence'] = ' '
    unsupported = client.post('/api/matches', json={'job_id': 'java', 'student': candidate})
    assert unsupported.status_code == 200
    java = next(item for item in unsupported.json()['items'] if item['tag_id'] == 'java')
    assert java['status'] == 'satisfied'
    assert java in unsupported.json()['satisfied_items']


@pytest.mark.parametrize('body', [{'student': {'private-resume-secret': 'secret-text'}, 'job_id': 'java'}, {'student': {'skills': [{'tag_id': 'java', 'label': 'Java', 'level': 99}]}, 'job_id': 'java'}])
def test_invalid_inputs_are_safe(client, body):
    r = client.post('/api/matches', json=body)
    assert r.status_code == 422
    assert set(r.json()['error']) == {'code', 'message', 'retryable', 'request_id'}
    assert 'secret-text' not in r.text


def test_health_jobs_sources(client):
    health = client.get('/api/health')
    assert health.status_code == 200 and 'API_KEY' not in health.text
    assert 'no-store' in health.headers['cache-control']
    jobs = client.get('/api/jobs').json()['items']
    assert len(jobs) == 6
    source_id = jobs[0]['requirements'][0]['evidence'][0]['source_id']
    assert client.get('/api/sources/' + source_id).json()['row'] > 1
    assert client.get('/api/tags').json()['items']
    assert client.get('/api/jobs/missing').status_code == 404


def test_llm_config_api_never_returns_api_key(client, monkeypatch):
    monkeypatch.delenv('LLM_ADAPTER', raising=False)
    monkeypatch.setenv('LLM_BASE_URL', '')
    monkeypatch.setenv('LLM_MODEL', '')
    monkeypatch.setenv('LLM_API_KEY', '')
    empty = client.get('/api/llm/config')
    assert empty.status_code == 200
    assert empty.json()['configured'] is False
    assert empty.json()['adapter'] == 'chat-completions'
    assert 'LLM_API_KEY' not in empty.text

    saved = client.post('/api/llm/config', json={
        'provider': 'deepseek',
        'adapter': 'chat-completions',
        'base_url': 'https://api.deepseek.com',
        'model': 'deepseek-chat',
        'api_key': 'test-secret-key'
    })
    assert saved.status_code == 200
    assert saved.json()['provider'] == 'deepseek'
    assert saved.json()['adapter'] == 'chat-completions'
    assert saved.json()['configured'] is True
    assert saved.json()['has_api_key'] is True
    assert 'test-secret-key' not in saved.text

    retained = client.post('/api/llm/config', json={
        'provider': 'openai',
        'adapter': 'chat-completions',
        'base_url': 'https://api.deepseek.com',
        'model': 'gpt-4o-mini'
    })
    assert retained.status_code == 200
    assert retained.json()['adapter'] == 'chat-completions'
    assert retained.json()['has_api_key'] is True

    changed_without_key = client.post('/api/llm/config', json={
        'provider': 'openai',
        'adapter': 'chat-completions',
        'base_url': 'https://api.openai.com/v1',
        'model': 'gpt-4o-mini'
    })
    assert changed_without_key.status_code == 400
    assert changed_without_key.json()['error']['code'] == 'LLM_CONFIG'
    assert 'test-secret-key' not in changed_without_key.text

    changed_with_key = client.post('/api/llm/config', json={
        'provider': 'openai',
        'adapter': 'chat-completions',
        'base_url': 'https://api.openai.com/v1',
        'model': 'gpt-4o-mini',
        'api_key': 'replacement-secret-key'
    })
    assert changed_with_key.status_code == 200
    assert changed_with_key.json()['has_api_key'] is True
    assert 'replacement-secret-key' not in changed_with_key.text


def test_llm_config_allows_localhost_address_change_without_reentering_key(client, monkeypatch):
    monkeypatch.setenv('LLM_BASE_URL', 'http://127.0.0.1:9000')
    monkeypatch.setenv('LLM_MODEL', 'local-model')
    monkeypatch.setenv('LLM_API_KEY', 'local-secret-key')
    response = client.post('/api/llm/config', json={
        'base_url': 'http://localhost:9001',
        'model': 'local-model'
    })
    assert response.status_code == 200
    assert response.json()['has_api_key'] is True
    assert 'local-secret-key' not in response.text


def test_llm_config_rejects_non_local_host(client):
    response = client.get('/api/llm/config', headers={'host': 'example.invalid'})
    assert response.status_code == 403
    assert response.json()['error']['code'] == 'HOST_NOT_ALLOWED'
    response = client.post('/api/llm/config', headers={'host': 'example.invalid'}, json={
        'base_url': 'https://api.openai.com/v1',
        'model': 'test-model',
        'api_key': 'secret'
    })
    assert response.status_code == 403
    assert response.json()['error']['code'] == 'HOST_NOT_ALLOWED'
    assert 'secret' not in response.text


def test_llm_connection_test_is_explicit_and_returns_no_secret(client, monkeypatch):
    monkeypatch.setenv('LLM_MODEL', 'test-model')
    async def fake_call_json(instruction, payload):
        assert 'connectivity_test' in payload['purpose']
        return {'ok': True}
    monkeypatch.setattr(main, 'call_json', fake_call_json)
    response = client.post('/api/llm/test', json={})
    assert response.status_code == 200
    assert response.json() == {'connected': True, 'model': 'test-model'}
    assert 'API_KEY' not in response.text


def test_llm_config_rejects_insecure_remote_url(client):
    response = client.post('/api/llm/config', json={
        'provider': 'openai',
        'base_url': 'http://example.com/v1',
        'model': 'test-model',
        'api_key': 'secret'
    })
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'LLM_CONFIG'
    assert 'secret' not in response.text


def test_llm_config_rejects_unknown_adapter(client):
    response = client.post('/api/llm/config', json={
        'provider': 'openai',
        'adapter': 'unknown-adapter',
        'base_url': 'https://api.openai.com/v1',
        'model': 'gpt-4o-mini',
        'api_key': 'secret'
    })
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'LLM_CONFIG'
    assert 'secret' not in response.text


def test_calculation_and_recommendation_consistent(client):
    r = client.post('/api/matches', json={'job_id': 'java', 'student': student()}).json()
    assert 'gap_items' not in r
    assert 'gap' not in r
    assert all('gap_reason' not in item for item in r['items'])
    assert 'gap' not in r['preferred_summary']
    rec = client.post('/api/recommendations', json={'student': student()}).json()
    java = next(x for x in rec['items'] if x['job_id'] == 'java')
    assert 'gap_items' not in java['match']
    assert 'gap' not in java['match']
    assert all('gap_reason' not in item for item in java['match']['items'])
    assert 'gap' not in java['match']['preferred_summary']
    assert java['match']['basic'] == r['basic']
    assert java['match']['input_version'] == r['input_version']
    changed = student()
    changed['major'] = '计算机科学'
    again = client.post('/api/matches', json={'job_id': 'java', 'student': changed}).json()
    assert again['input_version'] != r['input_version']


def test_report_recalculates_and_exports_server_facts(client, monkeypatch):
    async def advice(s, j, m):
        return {'fit_evaluation': '测试用建议', 'learning_directions': ['SQL'], 'learning_steps': ['编写查询脚本']}
    monkeypatch.setattr(main, 'generate_advice', advice)
    r = client.post('/api/reports', json={'job_id': 'java', 'student': student()})
    assert r.status_code == 200
    report = r.json()
    m = report['match']
    assert f"{m['basic']:.1f}%" in report['export_text']
    assert m['input_version'] in report['export_text']
    assert '北京' in report['export_text'] and '编写查询脚本' in report['export_text']
    assert report['status'] == 'complete'
    assert client.post('/api/reports', json={'job_id': 'java', 'student': student(), 'basic': 100}).status_code == 422


def test_upload_stays_in_memory_and_releases(client, monkeypatch):
    import starlette.formparsers
    original = starlette.formparsers.SpooledTemporaryFile
    tracked = []
    def spy(*a, **kw):
        result = original(*a, **kw)
        tracked.append(result)
        return result
    monkeypatch.setattr(starlette.formparsers, 'SpooledTemporaryFile', spy)
    async def resume(text):
        return {'profile': StudentProfile().model_dump(), 'notice': 'test stub', 'mode': 'test'}
    monkeypatch.setattr(main, 'extract_resume', resume)
    # DOCX/PDF may be >1MB but extracted text is bounded. Use extract_text stub to
    # isolate multipart buffering, not to claim a valid two-megabyte TXT resume.
    monkeypatch.setattr(main, 'extract_text', lambda *a: 'Java课程项目实习经历')
    r = client.post('/api/resume/parse', files={'file': ('resume.pdf', b'a' * (2 * 1024 * 1024), 'application/pdf')})
    assert r.status_code == 200
    assert tracked and all(x.closed and not x._rolled for x in tracked)


def test_upload_boundary(client):
    r = client.post('/api/resume/parse', files={'file': ('large.txt', b'x' * (MAX_BYTES + 100000))})
    assert r.status_code == 413 and r.json()['error']['code'] == 'FILE_TOO_LARGE'


def test_export_uses_same_half_up_display_as_match():
    from backend.app.matching import match_student
    from backend.app.models import Ability
    tags = ['java', 'sql', 'html', 'css', 'javascript', 'vue', 'linux', 'mysql']
    job = {'id': 'rounding', 'name': 'Rounding', 'requirements': [
        {'tag_id': tag, 'dimension': 'skills', 'required_level': 2, 'required': True} for tag in tags]}
    s = StudentProfile(confirmed=True, skills=[Ability(tag_id=tag, label=tag, level=1 if i == 0 else 0, confirmed=True, evidence='course exercise') for i, tag in enumerate(tags)])
    m = match_student(s, job)
    assert m['enhanced'] == 12.5
    assert m['enhanced_display'] == 12.5
    report = main.export_report(s, job, m, {'fit_evaluation': 'test', 'learning_directions': [], 'learning_steps': []})
    assert '增强匹配度：12.5%' in report


def test_generic_failure_never_echoes_secret(client, monkeypatch, caplog):
    def crash(*a):
        raise RuntimeError('secret resume body API key')
    monkeypatch.setattr(main, 'match_student', crash)
    r = client.post('/api/matches', json={'job_id': 'java', 'student': student()})
    assert r.status_code == 500
    assert 'secret resume body' not in r.text and 'secret resume body' not in caplog.text
