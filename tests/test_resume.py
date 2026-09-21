from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from backend.app.llm import AIError
from backend.app.resume import extract_text, MAX_BYTES, MAX_CHARS


def pdf_bytes(text=None, password=None):
    writer = PdfWriter()
    page = writer.add_blank_page(595, 842)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 12 Tf 50 780 Td (' + text + ') Tj ET').encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-8-sig', 'gb18030'])
def test_txt_encoding(encoding):
    text = '软件工程专业，使用Java完成课程项目，负责沟通与团队协作。'
    assert extract_text('resume.txt', text.encode(encoding)) == text


def test_docx_with_table():
    doc = Document()
    doc.add_paragraph('软件工程专业，Java课程项目。')
    doc.add_table(1, 1).cell(0, 0).text = 'MySQL数据库实践'
    output = BytesIO()
    doc.save(output)
    text = extract_text('resume.docx', output.getvalue())
    assert 'Java课程项目' in text and 'MySQL数据库实践' in text


def test_text_pdf():
    assert 'Java SQL coursework' in extract_text('resume.pdf', pdf_bytes('Java SQL coursework'))


@pytest.mark.parametrize('filename,content,code', [
    ('resume.txt', b'', 'FILE_EMPTY'),
    ('resume.exe', b'hello world!', 'FILE_TYPE'),
    ('resume.pdf', b'not a valid PDF', 'FILE_UNREADABLE'),
    ('resume.docx', b'not a valid ZIP', 'FILE_UNREADABLE'),
    ('resume.txt', b'abc\x00defabcdefgh', 'FILE_UNREADABLE'),
    ('resume.txt', b'\xff', 'FILE_UNREADABLE'),
    ('resume.txt', b'x' * (MAX_BYTES + 1), 'FILE_TOO_LARGE'),
    ('resume.txt', b'x' * (MAX_CHARS + 1), 'FILE_TOO_LONG'),
], ids=["empty", "type", "corrupt-pdf", "corrupt-docx", "binary-txt", "invalid-encoding", "oversized-bytes", "oversized-text"])
def test_bad_files(filename, content, code):
    with pytest.raises(AIError) as e:
        extract_text(filename, content)
    assert e.value.code == code


def test_no_text_scan_feedback():
    with pytest.raises(AIError) as e:
        extract_text('scan.pdf', pdf_bytes())
    assert e.value.code == 'FILE_NO_TEXT'


def test_encrypted_pdf():
    with pytest.raises(AIError) as e:
        extract_text('locked.pdf', pdf_bytes('Java SQL coursework', 'test-password'))
    assert e.value.code == 'FILE_ENCRYPTED'


def test_docx_expansion_limit():
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', b'x' * (11 * 1024 * 1024))
    with pytest.raises(AIError) as e:
        extract_text('huge.docx', output.getvalue())
    assert e.value.code == 'FILE_TOO_LARGE'
